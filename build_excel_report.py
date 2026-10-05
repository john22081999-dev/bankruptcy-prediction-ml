""" 
Κατα την υλοποιηση της εργασιας υπηρξε βοηθεια AI
"""
import pandas as pd
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

SRC_CSV = "outputs/balancedDataOutcomes.csv"
OUT_XLSX = "outputs/balancedDataOutcomes.xlsx"

df = pd.read_csv(SRC_CSV)

wb = openpyxl.Workbook()

ws = wb.active
ws.title = "Results"

extra_cols = ["Accuracy_formula", "Precision_formula", "Recall_formula",
              "F1_formula", "Specificity_Metric1", "MCC_Metric2"]
headers = list(df.columns) + extra_cols
for j, h in enumerate(headers, start=1):
    c = ws.cell(row=1, column=j, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="264653")
    c.alignment = Alignment(horizontal="center", wrap_text=True)

col_idx = {h: i + 1 for i, h in enumerate(df.columns)}
TP_L, TN_L, FP_L, FN_L = (get_column_letter(col_idx[c]) for c in ["TP", "TN", "FP", "FN"])

n_base = len(df.columns)
ACC_L = get_column_letter(n_base + 1)
PREC_L = get_column_letter(n_base + 2)
REC_L = get_column_letter(n_base + 3)
F1_L = get_column_letter(n_base + 4)
SPEC_L = get_column_letter(n_base + 5)
MCC_L = get_column_letter(n_base + 6)

for i, row in df.iterrows():
    r = i + 2
    for j, col in enumerate(df.columns, start=1):
        ws.cell(row=r, column=j, value=row[col])

    tp, tn, fp, fn = f"{TP_L}{r}", f"{TN_L}{r}", f"{FP_L}{r}", f"{FN_L}{r}"
    prec_cell, rec_cell = f"{PREC_L}{r}", f"{REC_L}{r}"

    ws.cell(row=r, column=n_base + 1, value=f"=IFERROR(({tp}+{tn})/({tp}+{tn}+{fp}+{fn}),0)")
    ws.cell(row=r, column=n_base + 2, value=f"=IFERROR({tp}/({tp}+{fp}),0)")
    ws.cell(row=r, column=n_base + 3, value=f"=IFERROR({tp}/({tp}+{fn}),0)")
    ws.cell(row=r, column=n_base + 4, value=f"=IFERROR(2*{prec_cell}*{rec_cell}/({prec_cell}+{rec_cell}),0)")
    ws.cell(row=r, column=n_base + 5, value=f"=IFERROR({tn}/({tn}+{fp}),0)")
    ws.cell(row=r, column=n_base + 6,
            value=f"=IFERROR(({tp}*{tn}-{fp}*{fn})/SQRT(({tp}+{fp})*({tp}+{fn})*({tn}+{fp})*({tn}+{fn})),0)")

for col_letter in [ACC_L, PREC_L, REC_L, F1_L, SPEC_L, MCC_L]:
    for r in range(2, len(df) + 2):
        ws[f"{col_letter}{r}"].number_format = "0.00"

# widths
for j in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(j)].width = 14
ws.freeze_panes = "A2"


ws2 = wb.create_sheet("PivotSummary")
ws2["A1"] = "Average TEST-set metrics per classifier and balancing regime (across all 4 folds)"
ws2["A1"].font = Font(bold=True, size=12)

classifiers = sorted(df["Classifier"].unique())
balances = ["unbalanced", "balanced"]

headers2 = ["Classifier", "Balance", "Avg_Accuracy", "Avg_Precision", "Avg_Recall", "Avg_F1", "Avg_ROC_AUC"]
for j, h in enumerate(headers2, start=1):
    c = ws2.cell(row=3, column=j, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="2a9d8f")

set_col = get_column_letter(col_idx["Set"])
class_col = get_column_letter(col_idx["Classifier"])
bal_col = get_column_letter(col_idx["Balance"])
acc_col = get_column_letter(col_idx["Accuracy"])
prec_col = get_column_letter(col_idx["Precision"])
rec_col = get_column_letter(col_idx["Recall"])
f1_col = get_column_letter(col_idx["F1"])
auc_col = get_column_letter(col_idx["ROC_AUC"])
n_rows = len(df) + 1  

row_ptr = 4
for clf in classifiers:
    for bal in balances:
        ws2.cell(row=row_ptr, column=1, value=clf)
        ws2.cell(row=row_ptr, column=2, value=bal)
        crit_common = (
            f'Results!${class_col}$2:${class_col}${n_rows+1}, $A{row_ptr}, '
            f'Results!${bal_col}$2:${bal_col}${n_rows+1}, $B{row_ptr}, '
            f'Results!${set_col}$2:${set_col}${n_rows+1}, "test"'
        )
        ws2.cell(row=row_ptr, column=3, value=f'=AVERAGEIFS(Results!${acc_col}$2:${acc_col}${n_rows+1}, {crit_common})')
        ws2.cell(row=row_ptr, column=4, value=f'=AVERAGEIFS(Results!${prec_col}$2:${prec_col}${n_rows+1}, {crit_common})')
        ws2.cell(row=row_ptr, column=5, value=f'=AVERAGEIFS(Results!${rec_col}$2:${rec_col}${n_rows+1}, {crit_common})')
        ws2.cell(row=row_ptr, column=6, value=f'=AVERAGEIFS(Results!${f1_col}$2:${f1_col}${n_rows+1}, {crit_common})')
        ws2.cell(row=row_ptr, column=7, value=f'=AVERAGEIFS(Results!${auc_col}$2:${auc_col}${n_rows+1}, {crit_common})')
        for col in range(3, 8):
            ws2.cell(row=row_ptr, column=col).number_format = "0.00"
        row_ptr += 1

last_row = row_ptr - 1
for j in range(1, 8):
    ws2.column_dimensions[get_column_letter(j)].width = 16


rule = ColorScaleRule(start_type="min", start_color="F8696B",
                       mid_type="percentile", mid_value=50, mid_color="FFEB84",
                       end_type="max", end_color="63BE7B")
ws2.conditional_formatting.add(f"F4:F{last_row}", rule)


chart = BarChart()
chart.type = "col"
chart.grouping = "clustered"
chart.title = "Average TEST F1 score by classifier and balancing regime"
chart.y_axis.title = "F1 score"
chart.x_axis.title = "Classifier"
chart.height = 10
chart.width = 24


ws2["A50"] = "Classifier"
ws2["B50"] = "unbalanced_F1"
ws2["C50"] = "balanced_F1"
for i, clf in enumerate(classifiers):
    r = 51 + i
    ws2.cell(row=r, column=1, value=clf)
    ws2.cell(row=r, column=2,
             value=f'=AVERAGEIFS(Results!${f1_col}$2:${f1_col}${n_rows+1}, '
                    f'Results!${class_col}$2:${class_col}${n_rows+1}, $A{r}, '
                    f'Results!${bal_col}$2:${bal_col}${n_rows+1}, "unbalanced", '
                    f'Results!${set_col}$2:${set_col}${n_rows+1}, "test")')
    ws2.cell(row=r, column=3,
             value=f'=AVERAGEIFS(Results!${f1_col}$2:${f1_col}${n_rows+1}, '
                    f'Results!${class_col}$2:${class_col}${n_rows+1}, $A{r}, '
                    f'Results!${bal_col}$2:${bal_col}${n_rows+1}, "balanced", '
                    f'Results!${set_col}$2:${set_col}${n_rows+1}, "test")')
    ws2.cell(row=r, column=2).number_format = "0.00"
    ws2.cell(row=r, column=3).number_format = "0.00"

data = Reference(ws2, min_col=2, max_col=3, min_row=50, max_row=50 + len(classifiers))
cats = Reference(ws2, min_col=1, min_row=51, max_row=50 + len(classifiers))
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
ws2.add_chart(chart, "A9")

wb.save(OUT_XLSX)
print("Saved:", OUT_XLSX)
